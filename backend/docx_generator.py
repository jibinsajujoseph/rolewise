from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
from models import ResumeContent

def generate_resume_docx(resume: ResumeContent, accepted_summary: str = None, accepted_bullets: dict = None) -> BytesIO:
    document = Document()
    
    # Configure styling for ATS friendly template
    style = document.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    
    # Contact info
    contact_p = document.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_run = contact_p.add_run(resume.contact.name)
    contact_run.bold = True
    contact_run.font.size = Pt(16)
    
    contact_details = []
    if resume.contact.email:
        contact_details.append(resume.contact.email)
    if resume.contact.phone:
        contact_details.append(resume.contact.phone)
    if resume.contact.location:
        contact_details.append(resume.contact.location)
    for link in resume.contact.links:
        contact_details.append(link)
        
    if contact_details:
        details_p = document.add_paragraph()
        details_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        details_p.add_run(" | ".join(contact_details))
    
    # Summary
    summary_text = accepted_summary if accepted_summary else resume.summary
    if summary_text:
        document.add_heading('Summary', level=1)
        document.add_paragraph(summary_text)
    
    # Experience
    if resume.experience:
        document.add_heading('Experience', level=1)
        for exp in resume.experience:
            p = document.add_paragraph()
            p.add_run(f"{exp.title}").bold = True
            p.add_run(f" | {exp.company}")
            
            date_loc = []
            if exp.start_date or exp.end_date:
                date_str = f"{exp.start_date or ''} - {exp.end_date or ''}".strip(" -")
                date_loc.append(date_str)
            if exp.location:
                date_loc.append(exp.location)
            
            if date_loc:
                date_p = document.add_paragraph()
                date_p.add_run(" | ".join(date_loc)).italic = True
                
            for bullet in exp.bullets:
                bullet_text = bullet
                if accepted_bullets and bullet in accepted_bullets:
                    bullet_text = accepted_bullets[bullet]
                document.add_paragraph(bullet_text, style='List Bullet')
                
    # Education
    if resume.education:
        document.add_heading('Education', level=1)
        for edu in resume.education:
            p = document.add_paragraph()
            p.add_run(edu.institution).bold = True
            p.add_run(f" | {edu.degree}")
            if edu.date:
                p.add_run(f" | {edu.date}")

    # Projects
    if resume.projects:
        document.add_heading('Projects', level=1)
        for proj in resume.projects:
            p = document.add_paragraph()
            p.add_run(proj.name).bold = True
            if proj.date:
                p.add_run(f" | {proj.date}")
            if proj.link:
                p.add_run(f" | {proj.link}")
            
            for bullet in proj.bullets:
                bullet_text = bullet
                if accepted_bullets and bullet in accepted_bullets:
                    bullet_text = accepted_bullets[bullet]
                document.add_paragraph(bullet_text, style='List Bullet')
                
    # Skills
    if resume.skills:
        document.add_heading('Skills', level=1)
        for skill_cat in resume.skills:
            p = document.add_paragraph()
            p.add_run(f"{skill_cat.category}: ").bold = True
            p.add_run(", ".join(skill_cat.items))
            
    # Certifications
    if resume.certifications:
        document.add_heading('Certifications', level=1)
        for cert in resume.certifications:
            document.add_paragraph(cert, style='List Bullet')
            
    # Save to buffer
    f = BytesIO()
    document.save(f)
    f.seek(0)
    return f
